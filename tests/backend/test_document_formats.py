"""One bounded round-trip / genuine OCR / authenticated atomic-import check."""
import base64
import io
import sys
from pathlib import Path
import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
sys.path.insert(0, str(Path(__file__).resolve().parents[2]/'apps'/'api'))
from main import app
from app.core import database
from app.core.security import create_access_token
from app.models.production import Production
from app.models.shot import Shot, Panel
from app.models.asset import Asset
from app.models.user import User, Role
from app.services import panel_media_service, import_service, document_export
from app.services.document_export import render_document, COLUMNS
from app.services.document_import import parse_document
from app.services.importer import map_headers
from PIL import Image, ImageDraw, ImageFont
from openpyxl import load_workbook
from pypdf import PdfReader

@pytest.mark.asyncio
async def test_formats_ocr_permission_images_raw_columns_and_rollback(tmp_path, monkeypatch):
    engine=create_async_engine(f"sqlite+aiosqlite:///{tmp_path/'documents.db'}")
    sessions=async_sessionmaker(engine,expire_on_commit=False,autoflush=False)
    async with engine.begin() as conn: await conn.run_sync(database.Base.metadata.create_all)
    async with sessions() as db:
        writer=Role(name='writer',permissions={'*':True}); reader=Role(name='reader',permissions={'production.read':True})
        db.add_all([User(id='writer',email='writer@documents.invalid',password_hash='unused',role=writer), User(id='reader',email='reader@documents.invalid',password_hash='unused',role=reader),Production(id='prod',name='合成制作测试',fps_num=24,fps_den=1)])
        await db.commit()
    media=tmp_path/'media'
    monkeypatch.setattr(database,'AsyncSessionLocal',sessions)
    for module in [panel_media_service,import_service,document_export]: monkeypatch.setattr(module,'MEDIA_ROOT',media)
    picture=Image.new('RGB',(1000,400),'white')
    ImageDraw.Draw(picture).text((30,60),'SHOT 001\nSynthetic OCR 123',font=ImageFont.load_default(size=48),fill='black')
    png=io.BytesIO(); picture.save(png,format='PNG')
    records=[({'display_number':'001','name':'合成测试 English 123','description':'原始内容 =1+2','voice_over':'测试旁白','duration':3,'duration_frames':72},[png.getvalue(),png.getvalue()])]
    for fmt in ['xlsx','docx','pdf']:
        content=render_document('合成分镜测试',records,COLUMNS,fmt)
        parsed=parse_document(content,'synthetic.'+fmt)
        assert len(parsed['rows'])==2 and parsed['images']
        joined='\n'.join(parsed['rows'][1])
        assert '原始内容 =1+2' in joined and 'English 123' in joined
        if fmt in {'xlsx','docx'}: assert len(parsed['images'])==2
        if fmt=='xlsx':
            sheet=load_workbook(io.BytesIO(content)).active
            description_column=next(cell.column for cell in sheet[1] if cell.value == '画面描述')
            cell=sheet.cell(2, description_column)
            assert cell.data_type == 's' and cell.value == '原始内容 =1+2'
        if fmt=='pdf':
            assert '合成分镜测试' in PdfReader(io.BytesIO(content)).pages[0].extract_text()
    for ext in ['png','jpg']:
        out=io.BytesIO(); picture.save(out,format='PNG' if ext=='png' else 'JPEG')
        parsed=parse_document(out.getvalue(),'scan.'+ext)
        assert 'Synthetic OCR 123' in parsed['rows'][1][1] and parsed['warnings']
    # A real image-only PDF forces the OCR path, not the PDF text extractor.
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.utils import ImageReader
    scan=io.BytesIO(); canvas=Canvas(scan,pagesize=(1000,400)); canvas.drawImage(ImageReader(io.BytesIO(png.getvalue())),0,0,width=1000,height=400); canvas.save()
    assert not PdfReader(io.BytesIO(scan.getvalue())).pages[0].extract_text()
    parsed=parse_document(scan.getvalue(),'scan.pdf')
    assert 'Synthetic OCR 123' in parsed['rows'][1][6] and parsed['warnings']
    headers=['镜号','镜头标题','画面描述','未映射原文']
    mapping=map_headers(headers)
    attachment={'row_index':0,'filename':'synthetic.png','data_base64':base64.b64encode(png.getvalue()).decode()}
    payload={'headers':headers,'rows':[['001','合成导入','=1+2','全部原文保留']], 'mapping':mapping, 'images':[attachment,attachment]}
    auth={'Authorization':'Bearer '+create_access_token({'sub':'writer'})}
    readonly={'Authorization':'Bearer '+create_access_token({'sub':'reader'})}
    async with AsyncClient(transport=ASGITransport(app=app),base_url='http://test') as client:
        preview={'filename':'sample.csv','file_base64':base64.b64encode('镜号,画面描述\n001,合成内容'.encode()).decode()}
        assert (await client.post('/api/v1/productions/prod/import-preview',json=preview,headers=readonly)).status_code==403
        assert (await client.post('/api/v1/productions/prod/import-preview',json=preview,headers=auth)).status_code==200
        result=await client.post('/api/v1/productions/prod/import-commit',json=payload,headers=auth)
        assert result.status_code==201, result.text
        for fmt in ['xlsx','docx','pdf']:
            exported=await client.get('/api/v1/productions/prod/export/'+fmt,headers=auth)
            assert exported.status_code==200, exported.text[:500]
            assert parse_document(exported.content,'export.'+fmt)['images']
        assert (await client.get('/api/v1/productions/prod/export/png',headers=auth)).status_code==404
        before=set(media.iterdir())
        bad={**payload,'headers':headers+['时长'],'rows':[payload['rows'][0]+['3'],['002','失败行','','','-1']], 'mapping':{**mapping,'duration':{'col':4}}}
        result=await client.post('/api/v1/productions/prod/import-commit',json=bad,headers=auth)
        assert result.status_code==400, result.text
        assert set(media.iterdir())==before, 'Failed batch must not leave media files'
    async with sessions() as db:
        assert (await db.execute(select(func.count(Shot.id)))).scalar()==1
        assert (await db.execute(select(func.count(Asset.id)))).scalar()==2
        assert (await db.execute(select(func.count(Panel.id)))).scalar()==2
    await engine.dispose()
