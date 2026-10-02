"""Bounded local document parsing; every result goes through the existing preview."""
from __future__ import annotations
import base64
import io
import threading
import zipfile
from functools import lru_cache
from pathlib import Path
from tempfile import TemporaryDirectory
from PIL import Image, ImageOps
from app.services.importer import parse_csv
from app.services.legacy_import_adapter import legacy_module

MAX_FILE_BYTES = 40 * 1024 * 1024
MAX_ROWS = 10000
MAX_COLUMNS = 200
OCR_LOCK = threading.Lock()


def check_archive(content: bytes):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(e.file_size for e in entries) > 100 * 1024 * 1024 or any(e.flag_bits & 1 for e in entries):
            raise ValueError('压缩文档过大或已加密，请拆分或解密后导入。')


@lru_cache(maxsize=1)
def ocr_engine():
    import onnxruntime
    onnxruntime.disable_telemetry_events()
    from rapidocr_onnxruntime import RapidOCR
    return RapidOCR(intra_op_num_threads=2, inter_op_num_threads=1)


def recognize(data: bytes):
    # ponytail: one CPU engine per process; use the existing job queue for sustained OCR traffic.
    with OCR_LOCK:
        result, _ = ocr_engine()(data)
    if not result:
        return '', 0.0
    return '\n'.join(item[1] for item in result), round(min(float(item[2]) for item in result), 3)


def validate_image(data: bytes):
    if len(data) > 10 * 1024 * 1024: raise ValueError('单张图片不得超过 10 MB。')
    with Image.open(io.BytesIO(data)) as image:
        if image.format not in {'PNG','JPEG','GIF','WEBP'} or image.width * image.height > 20_000_000:
            raise ValueError('图片格式不支持或超过 2000 万像素。')
        image.load()
        return Image.MIME[image.format]


def image_bytes(data: bytes):
    validate_image(data)
    with Image.open(io.BytesIO(data)) as image:
        out = io.BytesIO()
        ImageOps.exif_transpose(image).convert('RGB').save(out, format='PNG')
        if len(out.getvalue()) > 10 * 1024 * 1024:
            raise ValueError('解码后的图片超过 10 MB，请缩小后导入。')
        return out.getvalue()


def attachment(index: int, filename: str, data: bytes):
    mime = validate_image(data)
    return {'row_index': index, 'filename': Path(filename).name, 'mime': mime, 'data_base64': base64.b64encode(data).decode()}


def parse_document(content: bytes, filename: str):
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError('文件为空或超过 40 MB。')
    suffix = Path(filename).suffix.lower()
    images, warnings = [], []
    if suffix == '.csv':
        rows = parse_csv(content)
    elif suffix == '.xlsx':
        check_archive(content)
        import re
        from xml.etree import ElementTree as ET
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            for name in archive.namelist():
                if not name.startswith('xl/worksheets/') or not name.endswith('.xml'): continue
                root = ET.fromstring(archive.read(name))
                for cell in root.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c'):
                    ref = cell.get('r','')
                    match = re.fullmatch(r'([A-Z]+)([0-9]+)',ref)
                    if match:
                        col=0
                        for char in match[1]: col=col*26+ord(char)-64
                        if col > MAX_COLUMNS or int(match[2]) > MAX_ROWS+1:
                            raise ValueError('工作表单元格超出 10000 行 / 200 列范围，请裁剪空白格式或拆分。')
        rows, embedded = legacy_module('import_parsing').parse_xlsx_package(content, include_image_data=True)
        for item in embedded:
            if item['data_row'] < 1:
                raise ValueError('图片位于表头，无法关联镜头，请移到对应数据行。')
            images.append(attachment(item['data_row'] - 1, item['filename'], base64.b64decode(item['data'], validate=True)))
    elif suffix == '.docx':
        check_archive(content)
        from docx import Document
        doc = Document(io.BytesIO(content))
        rows = []
        for table in doc.tables:
            # One field-label table exported per shot is parsed as one record.
            if table.rows and len(table.columns) == 2 and table.cell(0, 0).text == '字段':
                headers = [r.cells[0].text for r in table.rows[1:]]
                values = [r.cells[1].text for r in table.rows[1:]]
                if not rows: rows = [headers]
                if rows[0] != headers: raise ValueError('Word 的字段表头不一致，请分开导入。')
                record_index = len(rows) - 1
                rows.append(values)
                for blip in table._tbl.xpath('.//a:blip'):
                    rel = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    if rel in doc.part.related_parts:
                        part = doc.part.related_parts[rel]
                        images.append(attachment(record_index, Path(part.partname).name, part.blob))
            elif not rows:
                rows = [[cell.text for cell in row.cells] for row in table.rows]
                for row_index, row in enumerate(table.rows[1:]):
                    for blip in row._tr.xpath('.//a:blip'):
                        rel = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                        if rel in doc.part.related_parts:
                            part = doc.part.related_parts[rel]
                            images.append(attachment(row_index, Path(part.partname).name, part.blob))
            else:
                raise ValueError('Word 包含多张不同结构的表格，请分别导入以保留字段和图片关联。')
        if not rows:
            text = '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
            rows = [['镜头标题','画面描述'], [Path(filename).stem, text]] if text else []
            for part in doc.part.related_parts.values():
                if part.content_type.startswith('image/'):
                    images.append(attachment(0, Path(part.partname).name, part.blob))
    elif suffix in {'.jpg', '.jpeg', '.png'}:
        raw = image_bytes(content)
        text, confidence = recognize(raw)
        rows = [['镜头标题','画面描述'], [Path(filename).stem, text]]
        images = [attachment(0, filename, content)]
        warnings.append(f'OCR 最低文字置信度 {confidence:.0%}；图片作为一个镜头保留，请核对识别文字与字段映射。' if text else '未识别到文字，保留原图片；请核对镜头标题和字段。')
    elif suffix == '.pdf':
        from pypdf import PdfReader, PdfWriter
        import pypdfium2 as pdfium
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted: raise ValueError('请解密 PDF 后再导入。')
        if len(reader.pages) > 30: raise ValueError('PDF 单次最多 30 页，请拆分导入。')
        warnings.append('PDF 按 SHOT 标记或页面识别，请核对镜头边界、识别文字与字段映射；保留页面原图与原文。')
        rows = [list(legacy_module('import_parsing').PDF_IMPORT_HEADERS)]
        rendered = pdfium.PdfDocument(content)
        try:
            with TemporaryDirectory(prefix='frameforge-import-') as tmp:
                for index, page in enumerate(reader.pages):
                    width, height = rendered[index].get_size()
                    scale = min(2, (20_000_000 / max(1, width * height)) ** .5)
                    bitmap = rendered[index].render(scale=scale)
                    try:
                        buf = io.BytesIO(); bitmap.to_pil().save(buf, format='PNG'); raw = image_bytes(buf.getvalue())
                    finally: bitmap.close()
                    start = len(rows) - 1
                    if (page.extract_text() or '').strip():
                        writer = PdfWriter(); writer.add_page(page)
                        path = Path(tmp) / 'page.pdf'
                        with path.open('wb') as stream: writer.write(stream)
                        parsed, embedded = legacy_module('import_parsing').parse_pdf_storyboard(path, preserve_source=True)
                        added = [header for header in parsed[0] if header not in rows[0]]
                        rows[0].extend(added)
                        for existing in rows[1:]: existing.extend([''] * len(added))
                        for values in parsed[1:]:
                            by_header = dict(zip(parsed[0],values))
                            rows.append([by_header.get(header,'') for header in rows[0]])
                        if len(parsed) == 2: images.append(attachment(start, f'page-{index+1}.png', raw))
                        else:
                            images.extend(attachment(start + image['data_row'] - 1, image['filename'], image['raw']) for image in embedded)
                    else:
                        text, confidence = recognize(raw)
                        rows.append([f'{len(rows):03d}', f'PDF 第 {index+1} 页', '', '', '', '', text, '', '', ''] + ['']*(len(rows[0])-10))
                        images.append(attachment(start, f'page-{index+1}.png', raw))
                        warnings.append(f'第 {index+1} 页已 OCR（最低置信度 {confidence:.0%}），按单镜头保留完整页面；请核对并映射。' if text else f'第 {index+1} 页未识别到文字，仍保留完整页面图片；请核对并填写字段。')
        finally: rendered.close()
    else:
        raise ValueError('支持 XLSX、CSV、DOCX、PDF、JPG、PNG；旧版 XLS/DOC 请另存为新版。')
    if len(rows) < 2 or len(rows) > MAX_ROWS + 1 or any(len(r) > MAX_COLUMNS for r in rows):
        raise ValueError('文档为空或超过 10000 行 / 200 列，请拆分后导入。')
    if len(images) > 1000 or sum(len(image['data_base64']) for image in images) > 64 * 1024 * 1024:
        raise ValueError('文档图片过多或过大，请拆分后导入。')
    if any(len(cell) > 10000 for row in rows for cell in row):
        raise ValueError('单元格超过 10000 字符，请拆分内容。')
    if any(i['row_index'] >= len(rows)-1 for i in images):
        raise ValueError('图片不在数据行内，无法可靠关联镜头。')
    return {'rows': rows, 'images': images, 'warnings': warnings}
