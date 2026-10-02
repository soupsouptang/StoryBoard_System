"""Genuine XLSX/DOCX/PDF from the canonical production and asset records."""
from __future__ import annotations
import io
import csv
import json
from threading import Lock
from html import escape
from pathlib import Path
from functools import lru_cache
from starlette.concurrency import run_in_threadpool
from sqlalchemy import select
from sqlalchemy.orm import selectinload
from app.core.exceptions import DomainError, NotFoundError
from app.models.production import Production, Sequence
from app.models.shot import Shot
from app.models.asset import Asset
from app.models.field import ProjectColumn, ShotColumnValue
from app.services.panel_media_service import MEDIA_ROOT
from app.services.image_crop_service import ImageCropService
from app.services.timecode_format import frames_to_tc

from app.services.column_catalog import BUILTIN_BINDINGS, BUILTIN_KEYS
from app.services.custom_field_service import CustomFieldService

COLUMNS = [(key, label) for key, (label, kind, _, _) in BUILTIN_BINDINGS.items() if kind != 'pending' and key != 'panel_image']
MIMES = {'xlsx':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'pdf':'application/pdf', 'csv':'text/csv; charset=utf-8'}

def require_export_permission(user):
    permissions = getattr(getattr(user, 'role', None), 'permissions', None) or {}
    if not (permissions.get('*') or any(permissions.get(key) for key in ('production.read', 'production.write', 'export.create'))):
        raise DomainError('当前账号没有导出项目的权限', code='FORBIDDEN')


async def export_fields(db, production_id, user):
    require_export_permission(user)
    catalog = await CustomFieldService.column_catalog(db, production_id)
    return [{ 'column_id': row['instance']['id'] if row['instance'] else None, 'key': row['instance']['column_key'] if row['instance'] else row['catalog_key'],
              'label': row['instance']['label'] if row['instance'] else row['label'],
              'column_class': row['column_class'] }
            for row in catalog if row['binding_kind'] != 'pending' and (
                row['instance'] and row['instance']['state'] != 'removed'
                or not row['instance'] and row['catalog_key'] in BUILTIN_KEYS)]

async def export_document(db, production_id, fmt, user, selected_fields=None):
    if fmt not in MIMES: raise DomainError('不支持的导出格式', code='VALIDATION_ERROR')
    available = await export_fields(db, production_id, user)
    allowlist = {row['key']: row for row in available}
    selected = list(allowlist) if selected_fields is None else list(dict.fromkeys(selected_fields))
    if not selected or any(key not in allowlist for key in selected):
        raise DomainError('导出字段为空、已删除或不属于当前项目，请刷新字段列表。', code='INVALID_EXPORT_FIELDS')
    if fmt == 'csv' and selected == ['panel_image']:
        raise DomainError('CSV 不支持图片，请至少选择一个文字字段。', code='INVALID_EXPORT_FIELDS')
    include_image = 'panel_image' in selected and fmt != 'csv'
    prod = (await db.execute(select(Production).where(Production.id == production_id, Production.deleted_at.is_(None)))).scalar_one_or_none()
    if prod is None: raise NotFoundError('项目不存在')
    shots = list((await db.execute(select(Shot).options(selectinload(Shot.panels)).where(Shot.production_id == production_id, Shot.deleted_at.is_(None)).order_by(Shot.sort_index, Shot.id).limit(2001))).scalars())
    if len(shots) > 2000: raise DomainError('单次文档导出最多 2000 镜头，请拆分后导出。', code='VALIDATION_ERROR')
    fields = list((await db.execute(select(ProjectColumn).where(ProjectColumn.production_id == production_id, ProjectColumn.state == "active", ProjectColumn.binding_kind == "custom").order_by(ProjectColumn.sort_index, ProjectColumn.id))).scalars())
    values = {(value.shot_id,value.column_id): value.value for value in (await db.execute(select(ShotColumnValue).join(Shot, Shot.id == ShotColumnValue.shot_id).where(Shot.production_id == production_id, Shot.deleted_at.is_(None)))).scalars()}
    sequences = {s.id:s.name for s in (await db.execute(select(Sequence).where(Sequence.production_id == production_id))).scalars()}
    assets = {a.id:a for a in (await db.execute(select(Asset).where(Asset.production_id == production_id, Asset.deleted_at.is_(None)))).scalars()}
    fps = prod.fps_num / (prod.fps_den or 1)
    frame = prod.start_timecode_frames
    records = []
    root = MEDIA_ROOT.resolve()
    total_media = 0
    for shot in shots:
        record = {key:getattr(shot,key,None) for key,_ in COLUMNS}
        record.update(sequence_id=sequences.get(shot.sequence_id,''), tc_in=frames_to_tc(frame, fps, prod.drop_frame))
        frame += shot.duration_frames
        record['camera_movement'] = ' / '.join(str(v) for v in shot.camera_movement.values())
        record['primary_method'] = ' / '.join([shot.primary_method, *shot.secondary_methods])
        for field in fields: record[CustomFieldService._column_key(field)] = values.get((shot.id,field.id))
        media = []
        for panel in sorted(shot.panels,key=lambda p:(p.sort_index,p.id)) if include_image else []:
            asset = assets.get(panel.asset_id)
            if panel.deleted_at is not None or asset is None: continue
            content, _, _ = await ImageCropService.rendered(db, asset, root, owner_type="panel", owner_id=panel.id)
            data = content if isinstance(content, bytes) else await run_in_threadpool(content.read_bytes)
            total_media += len(data)
            if total_media > 64*1024*1024: raise DomainError('导出图片超过 64 MB，请拆分。',code='VALIDATION_ERROR')
            media.append(data)
        records.append((record,media))
    columns = [(key, allowlist[key]['label']) for key in selected if key != 'panel_image']
    return prod.name, await run_in_threadpool(render_document, prod.name, records, columns, fmt, include_image)


def text(value):
    if value is None: return ''
    return json.dumps(value, ensure_ascii=False) if isinstance(value,(dict,list)) else str(value)


def render_document(title, records, columns, fmt, include_image=True):
    output = io.BytesIO()
    if fmt == 'csv':
        stream = io.StringIO(newline=''); writer = csv.writer(stream)
        writer.writerow([label for _,label in columns])
        writer.writerows([[text(record.get(key)) for key,_ in columns] for record,_ in records])
        return stream.getvalue().encode('utf-8-sig')
    if fmt == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.drawing.image import Image as ExcelImage
        from openpyxl.utils import get_column_letter
        workbook = Workbook(); sheet = workbook.active; sheet.title='分镜制作'
        sheet.append([*(['分镜画面'] if include_image else []), *[label for _,label in columns]])
        sheet.freeze_panes = 'C2'; sheet.auto_filter.ref = f'A1:{get_column_letter(len(columns)+int(include_image))}{len(records)+1}'
        streams = []
        for index,(record,media) in enumerate(records,2):
            sheet.append([*([''] if include_image else []), *[text(record.get(key)) for key,_ in columns]])
            sheet.row_dimensions[index].height = max(90, len(media)*80)
            for offset, data in enumerate(media):
                stream=io.BytesIO(data); streams.append(stream)
                image=ExcelImage(stream); image.width=150; image.height=80
                # Multiple images are stacked within their owning shot row.
                from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
                from openpyxl.drawing.xdr import XDRPositiveSize2D
                from openpyxl.utils.units import pixels_to_EMU
                image.anchor=OneCellAnchor(_from=AnchorMarker(col=0,row=index-1,rowOff=pixels_to_EMU(offset*82)),ext=XDRPositiveSize2D(pixels_to_EMU(150),pixels_to_EMU(80)))
                sheet.add_image(image)
        for row in sheet:
            for cell in row:
                # Never execute source text beginning with '=' as an Excel formula.
                cell.data_type='s'; cell.font=Font(name='Satoshi',bold=cell.row==1)
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        for index in range(1,len(columns)+int(include_image)+1): sheet.column_dimensions[get_column_letter(index)].width=24
        workbook.save(output)
    elif fmt == 'docx':
        from docx import Document
        from docx.shared import Inches
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        document = Document(); style=document.styles['Normal']; style.font.name='Satoshi'
        fonts=OxmlElement('w:rFonts'); fonts.set(qn('w:eastAsia'),'Sarasa UI SC'); style.element.get_or_add_rPr().append(fonts)
        document.add_heading(title,0)
        for record,media in records:
            document.add_heading('SHOT '+text(record.get('display_number')),1) if any(key == 'display_number' for key,_ in columns) else None
            table=document.add_table(rows=1,cols=2); table.style='Table Grid'
            table.rows[0].cells[0].text='字段'; table.rows[0].cells[1].text='内容'
            if include_image:
                cells=table.add_row().cells; cells[0].text='分镜画面'
                for data in media: cells[1].paragraphs[0].add_run().add_picture(io.BytesIO(data),width=Inches(5))
            for key,label in columns:
                cells=table.add_row().cells; cells[0].text=label; cells[1].text=text(record.get(key))
            document.add_page_break()
        document.save(output)
    elif fmt == 'pdf':
        from reportlab.pdfbase import pdfmetrics
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Image, Spacer, PageBreak
        from reportlab.lib.pagesizes import A4
        fonts = pdf_fonts()
        style=ParagraphStyle('body',fontName='Satoshi',fontSize=10,leading=14,wordWrap='CJK')
        flow=[]
        for index,(record,media) in enumerate(records):
            if index: flow.append(PageBreak())
            flow.append(Paragraph(pdf_text(title, fonts),style))
            if any(key == 'display_number' for key,_ in columns): flow.append(Paragraph(pdf_text('SHOT '+text(record.get('display_number')), fonts),style))
            flow.append(Spacer(1,12))
            for data in media:
                image=Image(io.BytesIO(data)); scale=min(450/image.imageWidth,230/image.imageHeight)
                image.drawWidth=image.imageWidth*scale; image.drawHeight=image.imageHeight*scale
                flow.extend([image,Spacer(1,8)])
            for key,label in columns:
                value=text(record.get(key))
                if value: flow.append(Paragraph(pdf_text(label+'：'+value, fonts),style))
        if not flow: flow=[Paragraph(pdf_text(title+' · 暂无镜头', fonts),style)]
        SimpleDocTemplate(output,pagesize=A4).build(flow)
    else: raise ValueError('Unsupported export format')
    return output.getvalue()

@lru_cache(maxsize=1)
def pdf_fonts():
    """Decode the same licensed, checked-in fonts used by the web app, once."""
    from fontTools.ttLib import TTFont as FontFile
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    root = Path(__file__).resolve().parents[3] / 'web' / 'public' / 'fonts'
    fonts=[]
    for index, path in enumerate([root/'satoshi-regular.woff2', *sorted(root.glob('sarasa-*.woff2'))]):
        name = 'Satoshi' if index == 0 else f'Sarasa{index}'
        font=FontFile(path); cmap=set(font.getBestCmap()); font.flavor=None
        # Sarasa web subsets share their original PostScript name; ReportLab
        # must distinguish them to keep every glyph's Unicode mapping.
        for record in font['name'].names:
            if record.nameID == 6: record.string = name.encode(record.getEncoding())
        stream=io.BytesIO(); font.save(stream); stream.seek(0)
        pdfmetrics.registerFont(TTFont(name,stream))
        fonts.append((name,cmap))
    return fonts


def pdf_text(value, fonts):
    runs=[]; last=None; run=''
    for char in value:
        if char == '\n':
            if run: runs.append(f'<font name="{last}">{escape(run)}</font>'); run=''
            runs.append('<br/>'); last=None; continue
        name=next((name for name,cmap in fonts if ord(char) in cmap),'Satoshi')
        if name != last and run: runs.append(f'<font name="{last}">{escape(run)}</font>'); run=''
        run+=char; last=name
    if run: runs.append(f'<font name="{last}">{escape(run)}</font>')
    return ''.join(runs)


# ponytail: serialize native PDFium rendering; move to an export worker if throughput requires it.
_pdf_preview_lock = Lock()

def render_pdf_preview(content, page_index):
    """Rasterize the actual export bytes; browsers do not need a PDF plugin."""
    import pypdfium2 as pdfium
    with _pdf_preview_lock:
        document = pdfium.PdfDocument(content)
        try:
            count = len(document)
            if page_index >= count:
                raise DomainError('预览页不存在，请返回第一页。', code='INVALID_PREVIEW_PAGE')
            page = document[page_index]
            try:
                bitmap = page.render(scale=min(1.5, 1200 / max(page.get_size())))
                try:
                    output = io.BytesIO(); bitmap.to_pil().save(output, format='PNG')
                    return output.getvalue(), count
                finally: bitmap.close()
            finally: page.close()
        finally: document.close()
